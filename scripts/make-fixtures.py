"""Reproducible synthetic fixtures. No user documents are read."""
from pathlib import Path
from hashlib import sha256
import io, json, zipfile, datetime
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'tests' / 'fixtures'
OUT.mkdir(parents=True, exist_ok=True)
font = ImageFont.truetype('C:/Windows/Fonts/msyh.ttc', 36)
def picture(label, color):
    image = Image.new('RGB', (900, 560), '#ffffff')
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((35, 35, 865, 525), radius=25, fill=color)
    draw.text((80, 100), label, fill='#0f172a', font=font)
    draw.rectangle((80, 240, 350, 350), fill='#2563eb')
    draw.ellipse((460, 220, 630, 390), fill='#f97316')
    return image
image = picture('轻阅 · 图片与中文检查', '#dbeafe')
for ext in ['png', 'jpg', 'bmp']: image.save(OUT / ('图片样本.' + ext))
image.save(OUT / '多页样本.tif', save_all=True, append_images=[picture('第二页 · 多页 TIFF', '#dcfce7')], compression='tiff_deflate')
image.save(OUT / '扫描样本.pdf', save_all=True, append_images=[picture('第二页 · 扫描 PDF', '#dcfce7')])

document = Document()
document.styles['Normal'].font.name = 'Microsoft YaHei'
document.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
document.sections[0].header.paragraphs[0].text = '轻阅测试样本 · 页眉'
document.sections[0].footer.paragraphs[0].text = '页脚 · 所有内容由测试脚本生成'
document.add_heading('轻阅文档阅读测试', 0)
document.add_paragraph('中文、English、标点，编号 001234567890123456。')
document.add_paragraph('第一段用于核对文字完整性。第二段用于检查换行和字体。')
table = document.add_table(rows=4, cols=3)
table.style = 'Table Grid'
for row, values in zip(table.rows, [('项目','数量','备注'),('苹果','12','普通表格'),('香蕉','3','中文显示'),('合计','15','核对数字')]):
    for cell, value in zip(row.cells, values): cell.text = value
document.add_picture(str(OUT / '图片样本.png'), width=Inches(4.6))
document.add_page_break()
document.add_heading('第二页检查', 1)
document.add_paragraph('分页内容应在第二页。搜索词：小熊读书。')
document.save(OUT / '阅读样本.docx')
# A real OOXML template has its template content type, not just a changed suffix.
with zipfile.ZipFile(OUT / '阅读样本.docx') as source, zipfile.ZipFile(OUT / '模板样本.dotx', 'w', zipfile.ZIP_DEFLATED) as dest:
    for entry in source.infolist():
        payload=source.read(entry.filename)
        if entry.filename=='[Content_Types].xml': payload=payload.replace(b'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml',b'application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml')
        dest.writestr(entry.filename,payload)

book = Workbook()
sheet = book.active; sheet.title = '中文数据'
sheet.append(['编号','日期','金额','合计公式','备注'])
sheet.append(['001234567890123456', datetime.date(2026,10,12), 1234.5, '=C2+C3', '小熊读书'])
sheet.append(['00002', datetime.date(2026,10,13), 20.75, None, '长编号与前导零'])
sheet['A5']='合并单元格'; sheet.merge_cells('A5:C5')
sheet['B2'].number_format=sheet['B3'].number_format='yyyy-mm-dd';sheet['C2'].number_format=sheet['C3'].number_format='#,##0.00';sheet['D2'].number_format='#,##0.00'
for cell in sheet[1]: cell.fill=PatternFill('solid',fgColor='DBEAFE');cell.font=Font(bold=True)
second=book.create_sheet('第二张表');second.append(['城市','人数']);second.append(['上海',42]);second.append(['北京',37])
large=book.create_sheet('分页搜索');large['A1']='开始';large['A501']='尾部搜索词';large['BC3']='右侧列测试'
book.save(OUT/'数据样本.xlsx')
with zipfile.ZipFile(OUT/'数据样本.xlsx') as source:
    payloads={entry.filename:source.read(entry.filename) for entry in source.infolist()}
payloads['xl/worksheets/sheet1.xml']=payloads['xl/worksheets/sheet1.xml'].replace(b'<f>C2+C3</f><v></v>',b'<f>C2+C3</f><v>1255.25</v>')
with zipfile.ZipFile(OUT/'数据样本.xlsx','w',zipfile.ZIP_DEFLATED) as dest:
    for name,payload in payloads.items():dest.writestr(name,payload)
with zipfile.ZipFile(OUT/'数据样本.xlsx') as source, zipfile.ZipFile(OUT/'表格模板.xltx','w',zipfile.ZIP_DEFLATED) as dest:
    for entry in source.infolist():
        payload=source.read(entry.filename)
        if entry.filename=='[Content_Types].xml':payload=payload.replace(b'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml',b'application/vnd.openxmlformats-officedocument.spreadsheetml.template.main+xml')
        dest.writestr(entry.filename,payload)
try:
    from pptx import Presentation
    from pptx.util import Inches as PptInches
    slides=Presentation()
    slide=slides.slides.add_slide(slides.slide_layouts[1]);slide.shapes.title.text='轻阅幻灯片测试';slide.placeholders[1].text='中文、文字和图片\n第 1 页'
    slide.shapes.add_picture(str(OUT/'图片样本.png'),PptInches(5),PptInches(2.6),width=PptInches(3.5))
    slide=slides.slides.add_slide(slides.slide_layouts[1]);slide.shapes.title.text='第二页检查';slide.placeholders[1].text='小熊读书\n幻灯片应有两页'
    slides.save(OUT/'演示样本.pptx')
except ImportError: pass

(OUT/'文本样本.txt').write_bytes('中文文本检查\r\n编号 001234567890123456\r\n搜索词：小熊读书\r\n'.encode('utf-8'))
(OUT/'数据样本.csv').write_text('编号,名称,金额\n001234567890123456,小熊读书,1234.50\n00002,"有逗号,有引号",20.75\n',encoding='utf-8-sig')
(OUT/'网页样本.html').write_text('<!doctype html><html><body><h1>网页文档测试</h1><p>中文正文与普通表格。</p><table border="1"><tr><td>编号</td><td>001234</td></tr></table><script>window.badScriptRan=true</script><img src="https://example.com/should-not-load.png" onerror="window.badScriptRan=true"><a href="javascript:alert(1)">安全检查</a></body></html>',encoding='utf-8')
(OUT/'网页样本.htm').write_bytes((OUT/'网页样本.html').read_bytes())
archive=MIMEMultipart('related');archive['Subject']='Synthetic web archive';part=MIMEText('<html><body><h1>网页归档测试</h1><p>中文静态内容 001234</p></body></html>','html','utf-8');part['Content-Location']='https://fixture.local/index.html';archive.attach(part)
for ext in ['mht','mhtml']:(OUT/('归档样本.'+ext)).write_bytes(archive.as_bytes())
(OUT/'通用样本.xml').write_text('<?xml version="1.0" encoding="utf-8"?><资料><标题>中文 XML 检查</标题><编号>001234</编号></资料>',encoding='utf-8')
(OUT/'办公表格.xml').write_text('<?xml version="1.0"?><Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet"><Worksheet ss:Name="中文表"><Table><Row><Cell><Data ss:Type="String">办公 XML 测试</Data></Cell></Row><Row><Cell><Data ss:Type="Number">42</Data></Cell></Row></Table></Worksheet></Workbook>',encoding='utf-8')
rtf=''.join(('\\u'+str(ord(c) if ord(c)<32768 else ord(c)-65536)+'?') if ord(c)>127 else c for c in '中文 RTF 检查：小熊读书')
(OUT/'富文本样本.rtf').write_text('{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Microsoft YaHei;}}\\f0\\fs32 '+rtf+'\\par Second paragraph.}',encoding='ascii')

def variant(source, target, before, after):
    with zipfile.ZipFile(OUT/source) as incoming, zipfile.ZipFile(OUT/target,'w',zipfile.ZIP_DEFLATED) as outgoing:
        for entry in incoming.infolist():
            payload=incoming.read(entry.filename)
            if entry.filename=='[Content_Types].xml':payload=payload.replace(before.encode(),after.encode())
            outgoing.writestr(entry.filename,payload)
variant('阅读样本.docx','静态宏类型.docm','application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml','application/vnd.ms-word.document.macroEnabled.main+xml')
variant('阅读样本.docx','静态宏模板.dotm','application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml','application/vnd.ms-word.template.macroEnabledTemplate.main+xml')
variant('数据样本.xlsx','静态宏表格.xlsm','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml','application/vnd.ms-excel.sheet.macroEnabled.main+xml')
variant('数据样本.xlsx','静态宏表格模板.xltm','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml','application/vnd.ms-excel.template.macroEnabled.main+xml')
if (OUT/'演示样本.pptx').exists():
    for ext,kind in [('pptm','application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml'),('potx','application/vnd.openxmlformats-officedocument.presentationml.template.main+xml'),('potm','application/vnd.ms-powerpoint.template.macroEnabled.main+xml'),('ppsx','application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml'),('ppsm','application/vnd.ms-powerpoint.slideshow.macroEnabled.main+xml')]:variant('演示样本.pptx','演示变体.'+ext,'application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml',kind)

index=[]
for file in sorted(OUT.iterdir()):
    if file.is_file():index.append({'file':file.name,'sha256':sha256(file.read_bytes()).hexdigest(),'bytes':file.stat().st_size,'source':'scripts/make-fixtures.py','origin':'synthetic','checks':['Chinese text','fixture-specific structure']})
(ROOT/'data'/'fixture-manifest.json').write_text(json.dumps({'schemaVersion':1,'fixtures':index},ensure_ascii=False,indent=2),encoding='utf-8')
print(f'Generated {len(index)} fixtures')
