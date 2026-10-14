"""Generate bounded larger documents under tmp for lifecycle comparisons."""
from pathlib import Path
from docx import Document
from openpyxl import Workbook
import sys
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject

root=Path(__file__).resolve().parent.parent
output=root/'tmp'/'memory-fixtures'
output.mkdir(parents=True,exist_ok=True)
pdf=PdfWriter()
page=pdf.add_blank_page(200,14400)
font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):pdf._add_object(font)})})
content=DecodedStreamObject()
content.set_data('\n'.join(f'BT /F1 10 Tf 10 {y} Td (Tall page {y}) Tj ET' for y in range(100,14400,200)).encode('ascii'))
page[NameObject('/Contents')]=pdf._add_object(content)
pdf.write(output/'超长页面样本.pdf')
if '--pdf-only' in sys.argv:raise SystemExit(0)
doc=Document()
for page in range(30):
    if page:doc.add_page_break()
    doc.add_heading(f'内存阅读样本 第 {page+1} 页',1)
    for paragraph in range(12):doc.add_paragraph(f'中文阅读检查 {page+1}-{paragraph+1}。'+'文字与标点完整，Memory sample 001234567890。'*12)
    if page%5==0:doc.add_picture(str(root/'tests'/'fixtures'/'图片样本.png'),width=3657600)
doc.save(output/'内存文字样本.docx')
book=Workbook(write_only=True)
sheet=book.create_sheet('五万行数据')
sheet.append(['编号','名称']+[f'数值{i}' for i in range(10)])
for row in range(50000):sheet.append([f'{row:018d}',f'中文数据 {row}']+[row+i/10 for i in range(10)])
book.save(output/'内存表格样本.xlsx')
book=Workbook(write_only=True)
sheet=book.create_sheet('六十万条匹配')
for row in range(50000):sheet.append([f'检索样本 {row}:{column}' for column in range(12)])
book.save(output/'搜索表格样本.xlsx')
print('Generated DOCX with 30 explicit page-break sections and XLSX with 50,000 data rows under tmp/memory-fixtures')
